import json
import os
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path


def atomic_json(path: Path, payload: dict) -> None:
    payload["updated_at"] = datetime.now(timezone.utc).isoformat()
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(temporary, path)


def run_job(config_path: Path) -> None:
    config_path = config_path.resolve()
    config = json.loads(config_path.read_text(encoding="utf-8"))
    status_path = Path(config["status_path"])

    def progress(percent: int, message: str) -> None:
        try:
            current = json.loads(status_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            current = {}
        if current.get("status") == "cancelled":
            raise RuntimeError("งานนี้ถูกยกเลิกเพื่อเริ่มงานใหม่")
        atomic_json(status_path, {
            "status": "processing",
            "progress": max(1, min(int(percent), 99)),
            "message": message,
        })

    try:
        scripts_dir = Path(__file__).resolve().parent
        sys.path.insert(0, str(scripts_dir))
        from engine_redbook import start_analysis_pipeline

        progress(2, "รับไฟล์แล้ว กำลังเตรียมระบบวิเคราะห์")
        summary = start_analysis_pipeline(
            config["input_path"],
            config["output_path"],
            product_desc=config["product_desc"],
            score_threshold=float(config["score_threshold"]),
            max_ai_rows=int(config["max_ai_rows"]),
            ollama_model=config["ollama_model"],
            use_ollama=bool(config["use_ollama"]),
            ollama_timeout=int(config["ollama_timeout"]),
            progress_callback=progress,
            embedding_cache_path=config.get("embedding_cache_path"),
        )

        from openpyxl import load_workbook

        workbook = load_workbook(config["output_path"], read_only=True, data_only=True)
        first_sheet = workbook[workbook.sheetnames[0]]
        preview = [list(row) for row in first_sheet.iter_rows(max_row=50, values_only=True)]
        sheet_names = list(workbook.sheetnames)
        workbook.close()

        atomic_json(status_path, {
            "status": "completed",
            "progress": 100,
            "message": "วิเคราะห์เสร็จสมบูรณ์",
            "result": {
                "summary": summary,
                "sheet_names": sheet_names,
                "preview": preview,
            },
        })
    except Exception as exception:
        error_path = status_path.with_name("worker-error.log")
        error_path.write_text(traceback.format_exc(), encoding="utf-8")
        current_status = {}
        try:
            current_status = json.loads(status_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            pass
        was_cancelled = current_status.get("status") == "cancelled"
        atomic_json(status_path, {
            "status": "cancelled" if was_cancelled else "failed",
            "progress": int(current_status.get("progress", 100)),
            "message": "งานถูกยกเลิก" if was_cancelled else "วิเคราะห์ไม่สำเร็จ",
            "error": str(exception)[:2000],
        })
        raise


def main() -> None:
    run_job(Path(sys.argv[1]))


if __name__ == "__main__":
    main()
