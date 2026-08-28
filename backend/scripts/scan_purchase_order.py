import argparse
import json
import mimetypes
import os
import tomllib
from pathlib import Path
from typing import Any

from google import genai
from google.genai import types
from openpyxl import load_workbook


PROMPT = """
คุณคือระบบอ่าน Purchase Order (PO) ของโรงพยาบาลและบริษัทในประเทศไทย
สกัดเฉพาะข้อมูลที่ปรากฏจริง ห้ามเดา ห้ามคำนวณข้อมูลที่เอกสารไม่ได้ระบุ
ตอบ JSON เท่านั้นโดยใช้โครงสร้างนี้:
{
  "po_number": "",
  "po_date": "YYYY-MM-DD หรือค่าว่าง",
  "customer_confirmed_at": "YYYY-MM-DD หรือค่าว่าง",
  "buyer_name": "",
  "tax_id": "",
  "address": "",
  "province": "",
  "district": "",
  "sub_district": "",
  "zipcode": "",
  "contact_name": "",
  "contact_channel": "",
  "payment_terms": "",
  "delivery_due_date": "YYYY-MM-DD หรือค่าว่าง",
  "delivery_terms": "",
  "quotation_valid_until": "YYYY-MM-DD หรือค่าว่าง",
  "note": "",
  "confidence": 0.0,
  "items": [
    {
      "source_code": "รหัสสินค้าใน PO",
      "source_name": "ชื่อสินค้าใน PO",
      "unit": "หน่วย",
      "quantity": 1,
      "unit_price": 0,
      "discount": 0,
      "confidence": 0.0
    }
  ],
  "warnings": []
}
วันที่ต้องเป็น ค.ศ. รูปแบบ YYYY-MM-DD ถ้าเอกสารเป็น พ.ศ. ให้ลบ 543
ตัวเลขต้องไม่มี comma และต้องเป็นตัวเลข หากหาไม่พบให้ใช้ 0 หรือค่าว่าง
confidence อยู่ระหว่าง 0 ถึง 1 และ warnings ระบุข้อความที่ Sale ควรตรวจสอบ
"""


def credentials() -> tuple[str, str]:
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    model = os.getenv("GEMINI_MODEL", "").strip()
    if not api_key:
        backend_root = Path(__file__).resolve().parents[1]
        with (backend_root / ".secrets" / "integrations.toml").open("rb") as file:
            secrets = tomllib.load(file)
        api_key = str(secrets["GEMINI_API_KEY"])
        model = model or str(secrets.get("GEMINI_MODEL", ""))
    return api_key, model or "gemini-3.5-flash"


def excel_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            lines.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(values_only=True):
                cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if cells:
                    lines.append(" | ".join(cells))
                if sum(len(line) for line in lines) > 30_000:
                    break
        return "\n".join(lines)
    finally:
        workbook.close()


def normalize(payload: dict[str, Any]) -> dict[str, Any]:
    def number(value: Any, default: float = 0.0) -> float:
        try:
            cleaned = str(value if value is not None else "").replace(",", "").replace("฿", "").strip()
            return float(cleaned) if cleaned else default
        except (TypeError, ValueError):
            return default

    items = payload.get("items") if isinstance(payload.get("items"), list) else []
    clean_items = []
    for item in items[:100]:
        if not isinstance(item, dict):
            continue
        clean_items.append({
            "source_code": str(item.get("source_code") or "").strip(),
            "source_name": str(item.get("source_name") or "").strip(),
            "unit": str(item.get("unit") or "").strip(),
            "quantity": max(1, int(number(item.get("quantity"), 1))),
            "unit_price": max(0.0, number(item.get("unit_price"))),
            "discount": max(0.0, number(item.get("discount"))),
            "confidence": min(1.0, max(0.0, number(item.get("confidence")))),
        })
    string_fields = [
        "po_number", "po_date", "customer_confirmed_at", "buyer_name", "tax_id", "address",
        "province", "district", "sub_district", "zipcode", "contact_name", "contact_channel",
        "payment_terms", "delivery_due_date", "delivery_terms", "quotation_valid_until", "note",
    ]
    result = {field: str(payload.get(field) or "").strip() for field in string_fields}
    result["confidence"] = min(1.0, max(0.0, number(payload.get("confidence"))))
    warnings = payload.get("warnings") if isinstance(payload.get("warnings"), list) else []
    result["warnings"] = [str(value).strip() for value in warnings if str(value).strip()][:20]
    result["items"] = clean_items
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document_path")
    args = parser.parse_args()
    path = Path(args.document_path)
    api_key, model = credentials()
    client = genai.Client(api_key=api_key)
    if path.suffix.lower() == ".xlsx":
        contents: list[Any] = [PROMPT, "ข้อมูลจาก Excel:\n" + excel_text(path)]
    else:
        mime_type = mimetypes.guess_type(path.name)[0] or "application/pdf"
        contents = [PROMPT, types.Part.from_bytes(data=path.read_bytes(), mime_type=mime_type)]
    response = client.models.generate_content(
        model=model,
        contents=contents,
        config=types.GenerateContentConfig(response_mime_type="application/json"),
    )
    clean = response.text.replace("```json", "").replace("```", "").strip()
    print(json.dumps(normalize(json.loads(clean)), ensure_ascii=False))


if __name__ == "__main__":
    main()
